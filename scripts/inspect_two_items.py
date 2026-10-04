import openpyxl
import psycopg2

db_url = 'postgresql://neondb_owner:npg_4B9AOcYnKbgu@ep-noisy-band-azabf3zx-pooler.c-3.ap-southeast-1.aws.neon.tech/neondb?sslmode=require'

wb = openpyxl.load_workbook(r'D:\movie\SADIA - INVENTORY FORMAT 1 og (3).xlsx', data_only=True, read_only=True)
sh = wb['DATA']
header = [c for c in next(sh.iter_rows(values_only=True))]

for r in sh.iter_rows(values_only=True):
    name = str(r[0]).strip() if r[0] else ''
    if name in ['Geepas Deep Fryer', 'Napkin Holder']:
        print(f"\n=== EXCEL DATA ROW: {name} ===")
        for h, v in zip(header, r):
            if v is not None:
                print(f"  {h:<25}: {v}")
wb.close()

conn = psycopg2.connect(db_url)
cur = conn.cursor()

for name in ['Geepas Deep Fryer', 'Napkin Holder']:
    print(f"\n=== DATABASE TRANSACTIONS: {name} ===")
    cur.execute("""
        SELECT t.id, t."transactionType", t.quantity, t."fromEntityType", t."fromEntityId", t."toEntityType", t."toEntityId", t.timestamp, t.notes
        FROM "InventoryTransaction" t
        JOIN "Product" p ON t."productId" = p.id
        WHERE p.name ILIKE %s AND t.timestamp < '2026-10-03 00:00:00'
        ORDER BY t.timestamp;
    """, (f"%{name}%",))
    rows = cur.fetchall()
    for row in rows:
        print(f"  {row[0]} | Type: {row[1]:<12} | Qty: {row[2]:5.1f} | From: {row[3]}:{row[4]} -> To: {row[5]}:{row[6]} | Date: {row[7]} | Notes: {row[8]}")

conn.close()
