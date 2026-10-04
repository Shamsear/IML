import sys
import psycopg2
import openpyxl
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

db_url = 'postgresql://neondb_owner:npg_4B9AOcYnKbgu@ep-noisy-band-azabf3zx-pooler.c-3.ap-southeast-1.aws.neon.tech/neondb?sslmode=require'
conn = psycopg2.connect(db_url)
cur = conn.cursor()

def pnum(v):
    try:
        if v is None or v == '': return 0.0
        return float(str(v).replace(',', '').strip())
    except:
        return 0.0

# 1. SOPEXA
sopexa_excel = {}
wb = openpyxl.load_workbook(r'D:\inventory\documents\Sopexa inventory - Master - 26-01-2026.xlsx', data_only=True, read_only=True)
sh = wb['INVENTORY']
for idx, r in enumerate(sh.iter_rows(values_only=True)):
    if idx < 4: continue
    if r[1] and str(r[1]).strip() and not str(r[1]).strip().upper().startswith('TOTAL'):
        name = str(r[1]).strip()
        sopexa_excel[name] = {
            'purchased': pnum(r[2]),
            'with_client': pnum(r[3]),
            'warehouse': pnum(r[4]),
            'rebrand': pnum(r[6]),
            'lost': pnum(r[7]),
            'used': pnum(r[8]),
            'damage': pnum(r[9]),
            'total': pnum(r[4])
        }
wb.close()

# 2. VIRGIN
virgin_excel = {}
wb = openpyxl.load_workbook(r'D:\inventory\documents\VIRGIN STOCK INVENTORY-2026.xlsx', data_only=True, read_only=True)
sh = wb['INVENTORY']
for idx, r in enumerate(sh.iter_rows(values_only=True)):
    if idx < 4: continue
    if r[1] and str(r[1]).strip() and not str(r[1]).strip().upper().startswith('TOTAL'):
        name = str(r[1]).strip()
        code = str(r[0]).strip() if r[0] else ''
        virgin_excel[name] = {
            'code': code,
            'purchased': pnum(r[2]),
            'issued': pnum(r[3]),
            'warehouse': pnum(r[4]),
            'damage': pnum(r[5]),
            'total': pnum(r[4])
        }
        if code:
            virgin_excel[code] = virgin_excel[name]
wb.close()

cur.execute("""
    SELECT p.id, p."itemCode", p.name, p.category, b.name as brand_name
    FROM "Product" p
    JOIN "Brand" b ON p."brandId" = b.id
    WHERE b.name NOT ILIKE '%SADIA%'
    ORDER BY b.name, p.name;
""")
other_prods = cur.fetchall()

cur.execute("""
    SELECT "productId", "transactionType", "fromEntityType", "toEntityType", SUM(quantity)
    FROM "InventoryTransaction"
    WHERE timestamp < '2026-10-03 00:00:00'
    GROUP BY "productId", "transactionType", "fromEntityType", "toEntityType";
""")
tx_aggs = cur.fetchall()

prod_txs = defaultdict(list)
for r in tx_aggs:
    prod_txs[r[0]].append({'type': r[1], 'from': r[2], 'to': r[3], 'qty': float(r[4])})

brand_counts = defaultdict(lambda: {'total': 0, 'checked': 0, 'matches': 0, 'discrepancies': []})

for p in other_prods:
    p_id, sku, name, cat, b_name = p
    brand_counts[b_name]['total'] += 1
    ex = {}
    if 'SOPEXA' in b_name.upper():
        clean_name = name
        if clean_name.lower().startswith('sopexa - '):
            clean_name = clean_name[9:].strip()
        ex = sopexa_excel.get(name) or sopexa_excel.get(clean_name) or {}
    elif 'VIRGIN' in b_name.upper():
        clean_name = name
        if clean_name.lower().startswith('virgin - '):
            clean_name = clean_name[9:].strip()
        ex = virgin_excel.get(name) or virgin_excel.get(clean_name) or virgin_excel.get(sku) or {}

    if not ex:
        continue

    brand_counts[b_name]['checked'] += 1
    txs = prod_txs[p_id]
    purchased = sum(t['qty'] for t in txs if t['type'] in ['RECEIVE', 'REC', 'INITIAL', 'INITIAL_STOCK', 'REBRAND_IN'])
    rebrand = sum(t['qty'] for t in txs if t['type'] in ['REBRAND', 'REBRAND_OUT'])
    used = sum(t['qty'] for t in txs if t['type'] == 'USED')
    issued = sum(t['qty'] for t in txs if t['type'] in ['ISSUE', 'OUT'] and t['to'] in ['STORE', None, ''])
    damage = sum(t['qty'] for t in txs if t['type'] in ['DAMAGE', 'DAM'])
    lost = sum(t['qty'] for t in txs if t['type'] in ['LOST', 'LST'])
    with_client = sum(t['qty'] for t in txs if t['type'] == 'CLIENT_STOCK' or (t['type'] in ['ISSUE', 'OUT'] and t['to'] in ['CLIENT', 'BRAND'])) - sum(t['qty'] for t in txs if t['type'] == 'CLIENT_RETURN')
    
    warehouse = max(0.0, purchased - issued - used - damage - lost - rebrand - with_client)

    diffs = {}
    for m, db_v in [('purchased', purchased), ('warehouse', warehouse), ('issued', issued), ('used', used), ('damage', damage), ('lost', lost), ('with_client', with_client), ('rebrand', rebrand)]:
        if m in ex:
            ex_v = ex.get(m, 0.0)
            if abs(db_v - ex_v) > 0.001:
                diffs[m] = (db_v, ex_v, db_v - ex_v)

    if diffs:
        brand_counts[b_name]['discrepancies'].append((sku, name, diffs))
    else:
        brand_counts[b_name]['matches'] += 1

print('=== OTHER BRANDS CROSS CHECK SUMMARY ===')
for b_name, stats in brand_counts.items():
    print(f'Brand: {b_name} | In DB: {stats["total"]} | Checked vs Excel: {stats["checked"]} | 100% Matches: {stats["matches"]} | Discrepancies: {len(stats["discrepancies"])}')
    for sku, name, diffs in stats['discrepancies'][:10]:
        print(f'   -- [{sku}] {name}')
        for m, (dv, ev, df) in diffs.items():
            print(f'        {m:<14}: DB={dv:8.1f} | Excel={ev:8.1f} | Diff={df:+8.1f}')

conn.close()
