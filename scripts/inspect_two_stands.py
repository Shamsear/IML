import openpyxl

wb_data = openpyxl.load_workbook(r'D:\movie\SADIA - INVENTORY FORMAT 1 og (3).xlsx', data_only=True, read_only=True)
sh = wb_data['DATA']
header = [c for c in next(sh.iter_rows(values_only=True))]

for r in sh.iter_rows(values_only=True):
    name = str(r[0]).strip() if r[0] else ''
    if name in ['Sadia Promotional DCC Stand  New 2025', 'Sadia Promotional Stand (1*1) - Generic New 2025 (Buy Scan & Win)']:
        print(f"\n=== ITEM: {name} ===")
        for h, v in zip(header, r):
            if v is not None:
                print(f"  {h:<25}: {v}")
wb_data.close()
